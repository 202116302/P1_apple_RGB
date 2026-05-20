"""
RGB 이미지 기반 과실 분석 플랫폼 - 학회 포스터용 흐름도 PPT 생성
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# 색상 정의
COLORS = {
    'primary': RGBColor(41, 128, 185),      # 파란색
    'secondary': RGBColor(39, 174, 96),     # 초록색
    'accent': RGBColor(230, 126, 34),       # 주황색
    'purple': RGBColor(142, 68, 173),       # 보라색
    'dark': RGBColor(44, 62, 80),           # 진한 회색
    'light': RGBColor(236, 240, 241),       # 밝은 회색
    'white': RGBColor(255, 255, 255),
    'red': RGBColor(231, 76, 60),
    'yellow': RGBColor(241, 196, 15),
}

def add_rounded_rectangle(slide, left, top, width, height, text, fill_color, font_size=14, bold=False):
    """둥근 사각형 추가"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top),
        Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = RGBColor(0, 0, 0)
    shape.line.width = Pt(1)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = COLORS['white'] if fill_color != COLORS['light'] and fill_color != COLORS['yellow'] else COLORS['dark']
    p.alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].space_before = Pt(0)
    tf.paragraphs[0].space_after = Pt(0)
    shape.text_frame.anchor = MSO_ANCHOR.MIDDLE

    return shape

def add_title_box(slide, left, top, width, height, text):
    """제목 박스 추가"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top),
        Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS['dark']
    shape.line.fill.background()

    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLORS['white']
    p.alignment = PP_ALIGN.CENTER
    shape.text_frame.anchor = MSO_ANCHOR.MIDDLE

    return shape

def create_slide_1_main_pipeline(prs):
    """슬라이드 1: 전체 시스템 파이프라인"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_title_box(slide, 0.5, 0.3, 12.33, 0.6, "RGB 이미지 기반 과실 분석 파이프라인")

    add_rounded_rectangle(slide, 0.7, 1.3, 1.8, 0.9, "Input Image\n(RGB)", COLORS['primary'], 12, True)

    slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(2.6), Inches(1.55), Inches(0.5), Inches(0.4)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = COLORS['dark']

    add_rounded_rectangle(slide, 3.2, 1.3, 2.2, 0.9, "Preprocessing\nWhite Balance\n+ CLAHE", COLORS['secondary'], 11)

    slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.5), Inches(1.55), Inches(0.5), Inches(0.4)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = COLORS['dark']

    add_rounded_rectangle(slide, 6.1, 1.3, 2.4, 0.9, "Object Detection\nYOLOv8n-seg\n(Instance Segment)", COLORS['accent'], 11)

    slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.6), Inches(1.55), Inches(0.5), Inches(0.4)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = COLORS['dark']

    add_rounded_rectangle(slide, 9.2, 1.3, 2.0, 0.9, "Cropped\nApple ROIs", COLORS['purple'], 12)

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(2.5), Inches(12.33), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS['dark']

    stage2_label = slide.shapes.add_textbox(Inches(0.5), Inches(2.6), Inches(3), Inches(0.4))
    tf = stage2_label.text_frame
    p = tf.paragraphs[0]
    p.text = "Stage 2: Quality Analysis (for each detected apple)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLORS['dark']

    add_rounded_rectangle(slide, 1.0, 3.3, 3.0, 1.2, "Disease Detection\nEfficientNet-B0\n(Binary Classification)", COLORS['red'], 12, True)
    add_rounded_rectangle(slide, 1.0, 4.8, 3.0, 0.7, "정상 / 질병", COLORS['light'], 12)
    slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(2.35), Inches(4.55), Inches(0.3), Inches(0.25)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = COLORS['dark']

    add_rounded_rectangle(slide, 4.7, 3.3, 3.0, 1.2, "Size Classification\nEfficientNet-B0\n(3-Class Classification)", COLORS['secondary'], 12, True)
    add_rounded_rectangle(slide, 4.7, 4.8, 3.0, 0.7, "L / M / S", COLORS['light'], 12)
    slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.05), Inches(4.55), Inches(0.3), Inches(0.25)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = COLORS['dark']

    add_rounded_rectangle(slide, 8.4, 3.3, 3.0, 1.2, "Weight Prediction\nEfficientNet-B0\n(Regression)", COLORS['primary'], 12, True)
    add_rounded_rectangle(slide, 8.4, 4.8, 3.0, 0.7, "무게 (g)", COLORS['light'], 12)
    slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.75), Inches(4.55), Inches(0.3), Inches(0.25)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = COLORS['dark']

    return slide

def create_slide_2_preprocessing(prs):
    """슬라이드 2: 이미지 전처리 과정"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_title_box(slide, 0.5, 0.3, 12.33, 0.6, "Image Preprocessing Pipeline")

    add_rounded_rectangle(slide, 1.0, 1.5, 2.5, 1.5, "Original\nRGB Image", COLORS['primary'], 14, True)

    slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(3.6), Inches(2.05), Inches(0.6), Inches(0.4)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = COLORS['dark']

    wb_box = add_rounded_rectangle(slide, 4.3, 1.2, 3.0, 2.1, "", COLORS['secondary'], 11)
    tf = wb_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "White Balance"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLORS['white']
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "\nGray World Algorithm"
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLORS['white']
    p2.alignment = PP_ALIGN.CENTER

    for formula in ["R' = R × (Avg / R_avg)", "G' = G × (Avg / G_avg)", "B' = B × (Avg / B_avg)"]:
        p = tf.add_paragraph()
        p.text = formula
        p.font.size = Pt(10)
        p.font.color.rgb = COLORS['light']
        p.alignment = PP_ALIGN.CENTER

    slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(7.4), Inches(2.05), Inches(0.6), Inches(0.4)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = COLORS['dark']

    clahe_box = add_rounded_rectangle(slide, 8.1, 1.2, 3.2, 2.1, "", COLORS['accent'], 11)
    tf = clahe_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "CLAHE"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLORS['white']
    p.alignment = PP_ALIGN.CENTER

    for text in ["\nContrast Limited Adaptive", "Histogram Equalization", "\n• LAB Color Space", "• clipLimit = 2.0", "• tileGridSize = 8×8"]:
        p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(10)
        p.font.color.rgb = COLORS['white'] if "•" not in text else COLORS['light']
        p.alignment = PP_ALIGN.CENTER

    purpose_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1.0), Inches(3.8),
        Inches(10.3), Inches(1.3)
    )
    purpose_box.fill.solid()
    purpose_box.fill.fore_color.rgb = COLORS['light']
    purpose_box.line.color.rgb = COLORS['dark']

    tf = purpose_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "전처리 목적"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLORS['dark']

    for text in ["• White Balance: 다양한 조명 조건에서 발생하는 색상 왜곡 보정", "• CLAHE: 이미지 대비 향상 및 밝기 편차 보정 (L채널만 적용하여 색상 왜곡 방지)"]:
        p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(11)
        p.font.color.rgb = COLORS['dark']

    return slide

def create_slide_3_results(prs):
    """슬라이드 3: 전체 시스템 성능 요약"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_title_box(slide, 0.5, 0.3, 12.33, 0.6, "Results: Model Performance Summary")

    summary_data = [
        ["Component", "Model", "Input Size", "Key Metric", "Performance", "Status"],
        ["Apple Detection", "YOLOv8n-seg", "640×640", "mAP@50", "99.44%", "Excellent"],
        ["Disease Detection", "EfficientNet-B0", "224×224", "Accuracy", "100.00%", "Excellent"],
        ["Size Classification", "EfficientNet-B0", "224×224", "Accuracy", "66.88%", "Needs Improvement"],
        ["Weight Prediction", "EfficientNet-B0", "224×224", "MAE", "30.11g", "Good"],
    ]

    table = slide.shapes.add_table(
        rows=5, cols=6,
        left=Inches(0.5), top=Inches(1.1),
        width=Inches(12.3), height=Inches(2.2)
    ).table

    status_colors = {
        "Excellent": COLORS['secondary'],
        "Good": COLORS['primary'],
        "Needs Improvement": COLORS['accent']
    }

    for i, row_data in enumerate(summary_data):
        for j, cell_data in enumerate(row_data):
            cell = table.cell(i, j)
            cell.text = cell_data
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.alignment = PP_ALIGN.CENTER
            cell.text_frame.anchor = MSO_ANCHOR.MIDDLE

            if i == 0:
                p.font.bold = True
                p.font.color.rgb = COLORS['white']
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS['dark']
            elif j == 5:
                p.font.bold = True
                p.font.color.rgb = COLORS['white']
                cell.fill.solid()
                cell.fill.fore_color.rgb = status_colors.get(cell_data, COLORS['dark'])
            elif j == 4:
                p.font.bold = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS['white']
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS['light'] if i % 2 == 0 else COLORS['white']

    t_title = slide.shapes.add_textbox(Inches(0.5), Inches(0.95), Inches(12.3), Inches(0.3))
    t_title.text_frame.paragraphs[0].text = "Table. Overall System Performance Summary"
    t_title.text_frame.paragraphs[0].font.size = Pt(12)
    t_title.text_frame.paragraphs[0].font.bold = True

    # 크기 등급 기준 테이블
    size_data = [
        ["Grade", "Korean Std.", "Weight Range", "Avg. Weight"],
        ["L", "특 (特)", "327 ~ 450 g", "363 g"],
        ["M", "상 (上)", "255 ~ 326 g", "282 g"],
        ["S", "보통", "185 ~ 254 g", "245 g"],
    ]

    table2 = slide.shapes.add_table(
        rows=4, cols=4,
        left=Inches(0.5), top=Inches(3.7),
        width=Inches(6.0), height=Inches(1.6)
    ).table

    row_colors = [COLORS['dark'], COLORS['red'], COLORS['secondary'], COLORS['primary']]

    for i, row_data in enumerate(size_data):
        for j, cell_data in enumerate(row_data):
            cell = table2.cell(i, j)
            cell.text = cell_data
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.alignment = PP_ALIGN.CENTER
            cell.text_frame.anchor = MSO_ANCHOR.MIDDLE

            if i == 0:
                p.font.bold = True
                p.font.color.rgb = COLORS['white']
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS['dark']
            elif j == 0:
                p.font.bold = True
                p.font.color.rgb = COLORS['white']
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_colors[i]
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS['light'] if i % 2 == 0 else COLORS['white']

    t2_title = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(6.0), Inches(0.3))
    t2_title.text_frame.paragraphs[0].text = "Size Grade Classification Criteria"
    t2_title.text_frame.paragraphs[0].font.size = Pt(12)
    t2_title.text_frame.paragraphs[0].font.bold = True

    # 데이터셋 정보
    dataset_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.8), Inches(3.5),
        Inches(6.0), Inches(1.8)
    )
    dataset_box.fill.solid()
    dataset_box.fill.fore_color.rgb = COLORS['primary']

    tf = dataset_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Dataset Information"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLORS['white']
    p.alignment = PP_ALIGN.CENTER

    for info in ["YOLO Segmentation: 76,856 images", "Size/Weight QC: 10,948 images", "Source: AI Hub 과실 품질 분류 데이터", "Variety: 후지 (Fuji) Apple"]:
        p = tf.add_paragraph()
        p.text = info
        p.font.size = Pt(11)
        p.font.color.rgb = COLORS['light']
        p.alignment = PP_ALIGN.CENTER

    return slide

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    create_slide_1_main_pipeline(prs)
    create_slide_2_preprocessing(prs)
    create_slide_3_results(prs)

    output_path = "docs/poster_methodology.pptx"
    import os
    os.makedirs("docs", exist_ok=True)
    prs.save(output_path)
    print(f"PowerPoint 생성 완료: {output_path}")

if __name__ == "__main__":
    main()
